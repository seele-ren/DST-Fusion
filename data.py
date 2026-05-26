from typing import Optional, Tuple
import math
import json

import numpy as np
import torch
from torch.utils.data import Dataset
import csv
import re
from datetime import datetime
from typing import Dict, List, Set


def compute_feature_stats(x: np.ndarray, eps: float = 1e-6) -> Tuple[np.ndarray, np.ndarray]:
    # 计算全样本、全时间的均值与标准差
    mean = x.mean(axis=(0, 1))
    std = x.std(axis=(0, 1))
    std = np.where(std < eps, 1.0, std)
    return mean, std


def compute_weekly_stats(weekly: Optional[np.ndarray], eps: float = 1e-6) -> Tuple[np.ndarray, np.ndarray]:
    if weekly is None or weekly.size == 0:
        return np.zeros(0, dtype=np.float32), np.ones(0, dtype=np.float32)
    mean = weekly.mean(axis=(0, 1))
    std = weekly.std(axis=(0, 1))
    std = np.where(std < eps, 1.0, std)
    return mean, std


def normalize_features(x: np.ndarray, mean: np.ndarray, std: np.ndarray) -> np.ndarray:
    return (x - mean[None, None, :]) / std[None, None, :]


def _is_peak_doy_col(name: str) -> bool:
    lowered = name.lower()
    return "peak_doy" in lowered or "peakdoy" in lowered


def expand_static_cols(static_cols: List[str]) -> List[str]:
    expanded: List[str] = []
    for col in static_cols:
        if _is_peak_doy_col(col):
            expanded.extend([f"{col}_sin", f"{col}_cos"])
        else:
            expanded.append(col)
    return expanded


def _iter_geo_points(geom):
    geom_type = geom.get("type")
    coords = geom.get("coordinates", [])
    if geom_type == "Polygon":
        for ring in coords:
            for lon, lat in ring:
                yield lon, lat
    elif geom_type == "MultiPolygon":
        for poly in coords:
            for ring in poly:
                for lon, lat in ring:
                    yield lon, lat


def build_county_adjacency(geojson_path: str, name_key: str = "name", rounding: int = 6):
    """Build adjacency matrix where counties touch at edges or vertices."""
    with open(geojson_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    features = data.get("features", [])
    county_names: List[str] = []
    point_to_counties: dict[tuple[float, float], List[int]] = {}

    for feat in features:
        props = feat.get("properties", {})
        name = props.get(name_key)
        if not name:
            continue
        county_idx = len(county_names)
        county_names.append(str(name))
        geom = feat.get("geometry", {})
        for lon, lat in _iter_geo_points(geom):
            key = (round(lon, rounding), round(lat, rounding))
            point_to_counties.setdefault(key, []).append(county_idx)

    n = len(county_names)
    adj = np.zeros((n, n), dtype=np.float32)
    for counties in point_to_counties.values():
        if len(counties) < 2:
            continue
        unique = list(set(counties))
        for i in range(len(unique)):
            for j in range(i + 1, len(unique)):
                a, b = unique[i], unique[j]
                adj[a, b] = 1.0
                adj[b, a] = 1.0
    return county_names, adj


def build_graph_year_samples(
    x: np.ndarray,
    y: np.ndarray,
    years: np.ndarray,
    counties: np.ndarray,
    county_order: List[str],
    weekly: Optional[np.ndarray] = None,
    static: Optional[np.ndarray] = None,
    province_targets: Optional[Dict[int, float]] = None,
) -> Tuple[
    List[Tuple[np.ndarray, np.ndarray, Optional[np.ndarray], Optional[np.ndarray], np.ndarray, float]],
    List[int],
]:
    """Assemble per-year graph samples (no normalization, no torch conversion)."""
    samples: List[
        Tuple[np.ndarray, np.ndarray, Optional[np.ndarray], Optional[np.ndarray], np.ndarray, float]
    ] = []
    sample_years: List[int] = []
    county_index = {name: i for i, name in enumerate(county_order)}
    unique_years = sorted(set(int(y) for y in years.tolist()))
    n_counties = len(county_order)
    seq_len = x.shape[1]
    feat_dim = x.shape[2]
    static_dim = static.shape[1] if static is not None else 0
    weekly_len = weekly.shape[1] if weekly is not None else 0
    weekly_dim = weekly.shape[2] if weekly is not None else 0

    for year in unique_years:
        x_year = np.zeros((n_counties, seq_len, feat_dim), dtype=np.float32)
        y_year = np.zeros((n_counties,), dtype=np.float32)
        static_year = (
            np.zeros((n_counties, static_dim), dtype=np.float32)
            if static is not None
            else None
        )
        weekly_year = (
            np.zeros((n_counties, weekly_len, weekly_dim), dtype=np.float32)
            if weekly is not None
            else None
        )
        mask = np.zeros((n_counties,), dtype=bool)
        year_mask = years == year
        indices = np.where(year_mask)[0]
        for idx in indices:
            cname = str(counties[idx])
            if cname not in county_index:
                continue
            cidx = county_index[cname]
            x_year[cidx] = x[idx]
            y_year[cidx] = y[idx]
            if static_year is not None:
                static_year[cidx] = static[idx]
            if weekly_year is not None:
                weekly_year[cidx] = weekly[idx]
            mask[cidx] = True
        if mask.any():
            if province_targets is not None:
                if int(year) not in province_targets:
                    raise ValueError(f"Missing province target for year {int(year)}")
                province_y = float(province_targets[int(year)])
            else:
                province_y = float(y_year[mask].mean())
            samples.append((x_year, y_year, static_year, weekly_year, mask, province_y))
            sample_years.append(year)
    return samples, sample_years


class GraphYearDataset(Dataset):
    """Build year-level graph samples with county adjacency."""

    def __init__(
        self,
        x: np.ndarray,
        y: np.ndarray,
        years: np.ndarray,
        counties: np.ndarray,
        county_order: List[str],
        adj: np.ndarray,
        weekly: Optional[np.ndarray] = None,
        static: Optional[np.ndarray] = None,
        mean: Optional[np.ndarray] = None,
        std: Optional[np.ndarray] = None,
        y_mean: Optional[float] = None,
        y_std: Optional[float] = None,
        weekly_mean: Optional[np.ndarray] = None,
        weekly_std: Optional[np.ndarray] = None,
        static_mean: Optional[np.ndarray] = None,
        static_std: Optional[np.ndarray] = None,
        province_targets: Optional[Dict[int, float]] = None,
        samples: Optional[
            List[
                Tuple[np.ndarray, np.ndarray, Optional[np.ndarray], Optional[np.ndarray], np.ndarray, float]
            ]
        ] = None,
        sample_years: Optional[List[int]] = None,
    ) -> None:
        self.mean = mean
        self.std = std
        self.y_mean = y_mean
        self.y_std = y_std
        self.weekly_mean = weekly_mean
        self.weekly_std = weekly_std
        self.static_mean = static_mean
        self.static_std = static_std
        self.adj = torch.from_numpy(adj.astype(np.float32))
        self.county_order = county_order
        self.county_index = {name: i for i, name in enumerate(county_order)}

        if samples is None or sample_years is None:
            self.samples = []
            self.sample_years = []
            unique_years = sorted(set(int(y) for y in years.tolist()))
            n_counties = len(county_order)
            seq_len = x.shape[1]
            feat_dim = x.shape[2]
            static_dim = static.shape[1] if static is not None else 0
            weekly_len = weekly.shape[1] if weekly is not None else 0
            weekly_dim = weekly.shape[2] if weekly is not None else 0

            for year in unique_years:
                x_year = np.zeros((n_counties, seq_len, feat_dim), dtype=np.float32)
                y_year = np.zeros((n_counties,), dtype=np.float32)
                static_year = (
                    np.zeros((n_counties, static_dim), dtype=np.float32)
                    if static is not None
                    else None
                )
                weekly_year = (
                    np.zeros((n_counties, weekly_len, weekly_dim), dtype=np.float32)
                    if weekly is not None
                    else None
                )
                mask = np.zeros((n_counties,), dtype=bool)
                year_mask = years == year
                indices = np.where(year_mask)[0]
                for idx in indices:
                    cname = str(counties[idx])
                    if cname not in self.county_index:
                        continue
                    cidx = self.county_index[cname]
                    x_year[cidx] = x[idx]
                    y_year[cidx] = y[idx]
                    if static_year is not None:
                        static_year[cidx] = static[idx]
                    if weekly_year is not None:
                        weekly_year[cidx] = weekly[idx]
                    mask[cidx] = True
                if mask.any():
                    if province_targets is not None:
                        if int(year) not in province_targets:
                            raise ValueError(f"Missing province target for year {int(year)}")
                        province_y = float(province_targets[int(year)])
                    else:
                        province_y = float(y_year[mask].mean())
                    self.samples.append((x_year, y_year, static_year, weekly_year, mask, province_y))
                    self.sample_years.append(year)
        else:
            self.samples = samples
            self.sample_years = sample_years

        if self.samples:
            normalized_samples = []
            for sample in self.samples:
                if len(sample) == 6:
                    x_year, y_year, static_year, weekly_year, mask, province_y = sample
                else:
                    x_year, y_year, static_year, weekly_year, mask = sample
                    province_y = float(y_year[mask].mean())
                if self.mean is not None and self.std is not None:
                    x_year = (x_year - self.mean[None, None, :]) / self.std[None, None, :]
                if self.y_mean is not None and self.y_std is not None:
                    y_year = (y_year - self.y_mean) / self.y_std
                    province_y = (province_y - float(self.y_mean)) / float(self.y_std)
                if static_year is not None and self.static_mean is not None and self.static_std is not None:
                    static_year = (static_year - self.static_mean[None, :]) / self.static_std[None, :]
                if weekly_year is not None and self.weekly_mean is not None and self.weekly_std is not None:
                    weekly_year = (weekly_year - self.weekly_mean[None, None, :]) / self.weekly_std[None, None, :]
                normalized_samples.append(
                    (
                        torch.from_numpy(x_year.astype(np.float32)),
                        torch.from_numpy(y_year.astype(np.float32)),
                        torch.from_numpy(static_year.astype(np.float32)) if static_year is not None else None,
                        torch.from_numpy(weekly_year.astype(np.float32)) if weekly_year is not None else None,
                        torch.from_numpy(mask.astype(np.bool_)),
                        torch.tensor(float(province_y), dtype=torch.float32),
                    )
                )
            self.samples = normalized_samples

    def __len__(self) -> int:
        return len(self.samples)

    def __getitem__(self, idx: int):
        x, y, static, weekly, mask, province_y = self.samples[idx]
        result = [x, y, mask]
        if weekly is not None:
            result.insert(2, weekly)
        if static is not None:
            insert_at = 3 if weekly is not None else 2
            result.insert(insert_at, static)
        result.append(province_y)
        result.append(self.adj)
        return tuple(result)


def _parse_date(date_str: str) -> datetime:
    # 支持多种日期格式
    for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%Y/%m/%d %H:%M:%S"):
        try:
            return datetime.strptime(date_str, fmt)
        except ValueError:
            continue
    raise ValueError(f"不支持的日期格式: {date_str}")


def _fill_missing(seq: np.ndarray, method: str) -> np.ndarray:
    if method == "zero":
        return np.nan_to_num(seq, nan=0.0)

    filled = seq.copy()
    # 前向填充
    for i in range(1, filled.shape[0]):
        missing = np.isnan(filled[i])
        if missing.any():
            filled[i, missing] = filled[i - 1, missing]
    # 处理开头的缺失值
    for i in range(filled.shape[0] - 2, -1, -1):
        missing = np.isnan(filled[i])
        if missing.any():
            filled[i, missing] = filled[i + 1, missing]
    return np.nan_to_num(filled, nan=0.0)


def _load_yield_table(
    path: str,
    yield_col: str,
    yield_county_col: str,
    yield_year_col: str,
) -> Dict[Tuple[str, int], float]:
    if path.lower().endswith(".csv"):
        # 读取 CSV 产量表
        with open(path, "r", newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            headers = reader.fieldnames or []
            for col in [yield_col, yield_county_col, yield_year_col]:
                if col not in headers:
                    raise ValueError(f"产量文件缺少列: {col}")
            targets: Dict[Tuple[str, int], float] = {}
            for row in reader:
                county = row[yield_county_col]
                year = int(row[yield_year_col])
                if row[yield_col] != "":
                    targets[(county, year)] = float(row[yield_col])
            return targets

    if path.lower().endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError("读取 .xlsx 需要安装 openpyxl") from exc

        # 读取 Excel 产量表
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise ValueError("产量文件为空")

        header_map = {str(h).strip(): idx for idx, h in enumerate(headers) if h is not None}
        for col in [yield_col, yield_county_col, yield_year_col]:
            if col not in header_map:
                raise ValueError(f"产量文件缺少列: {col}")

        targets = {}
        for row in rows:
            county = row[header_map[yield_county_col]]
            year = row[header_map[yield_year_col]]
            yval = row[header_map[yield_col]]
            if county is None or year is None or yval is None:
                continue
            targets[(str(county), int(year))] = float(yval)
        return targets

    raise ValueError("产量文件格式不支持，请使用 .csv 或 .xlsx")


def _parse_year_value(value) -> Optional[int]:
    if value is None:
        return None
    if isinstance(value, (int, np.integer)):
        y = int(value)
        return y if 1800 <= y <= 2200 else None
    if isinstance(value, (float, np.floating)):
        if np.isnan(value):
            return None
        y = int(value)
        return y if 1800 <= y <= 2200 else None
    text = str(value).strip()
    if not text:
        return None
    m = re.search(r"(18|19|20|21)\d{2}", text)
    return int(m.group(0)) if m else None


def load_province_year_targets(
    path: str,
    year_col: Optional[str] = None,
    target_col: Optional[str] = None,
    target_divisor: float = 1.0,
) -> Dict[int, float]:
    """Load province yearly targets from CSV/XLSX into {year: target}."""
    if not path:
        return {}

    divisor = float(target_divisor) if target_divisor else 1.0
    if abs(divisor) < 1e-12:
        divisor = 1.0

    if path.lower().endswith(".csv"):
        with open(path, "r", newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            headers = reader.fieldnames or []
            rows = list(reader)
        year_col_eff = year_col if year_col in headers else None
        target_col_eff = target_col if target_col in headers else None
        if year_col_eff is None:
            for col in headers:
                sample = [r.get(col) for r in rows[:20] if r.get(col) not in ("", None)]
                ok = sum(1 for v in sample if _parse_year_value(v) is not None)
                if sample and ok >= max(3, int(len(sample) * 0.6)):
                    year_col_eff = col
                    break
        if target_col_eff is None:
            best_col = None
            best_ratio = -1.0
            for col in headers:
                if col == year_col_eff:
                    continue
                vals = [r.get(col) for r in rows if r.get(col) not in ("", None)]
                if not vals:
                    continue
                ok = 0
                for v in vals:
                    try:
                        float(v)
                        ok += 1
                    except Exception:
                        pass
                ratio = ok / len(vals)
                if ratio > best_ratio:
                    best_ratio = ratio
                    best_col = col
            target_col_eff = best_col
        if year_col_eff is None or target_col_eff is None:
            raise ValueError(f"Cannot infer province cols from CSV headers: {headers}")
        out: Dict[int, float] = {}
        for r in rows:
            y = _parse_year_value(r.get(year_col_eff))
            if y is None:
                continue
            v = r.get(target_col_eff)
            if v in ("", None):
                continue
            try:
                out[y] = float(v) / divisor
            except Exception:
                continue
        return out

    if path.lower().endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError("读取 .xlsx 需要安装 openpyxl") from exc

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            return {}
        headers_s = [str(h).strip() if h is not None else "" for h in headers]
        data_rows = list(rows)
        year_col_eff = year_col if year_col in headers_s else None
        target_col_eff = target_col if target_col in headers_s else None
        if year_col_eff is None:
            for col in headers_s:
                idx = headers_s.index(col)
                sample = [r[idx] for r in data_rows[:20] if idx < len(r) and r[idx] is not None]
                ok = sum(1 for v in sample if _parse_year_value(v) is not None)
                if sample and ok >= max(3, int(len(sample) * 0.6)):
                    year_col_eff = col
                    break
        if target_col_eff is None:
            best_col = None
            best_ratio = -1.0
            for col in headers_s:
                if col == year_col_eff:
                    continue
                idx = headers_s.index(col)
                vals = [r[idx] for r in data_rows if idx < len(r) and r[idx] is not None]
                if not vals:
                    continue
                ok = 0
                for v in vals:
                    try:
                        float(v)
                        ok += 1
                    except Exception:
                        pass
                ratio = ok / len(vals)
                if ratio > best_ratio:
                    best_ratio = ratio
                    best_col = col
            target_col_eff = best_col
        if year_col_eff is None or target_col_eff is None:
            raise ValueError(f"Cannot infer province cols from XLSX headers: {headers_s}")
        y_idx = headers_s.index(year_col_eff)
        t_idx = headers_s.index(target_col_eff)
        out: Dict[int, float] = {}
        for r in data_rows:
            if y_idx >= len(r) or t_idx >= len(r):
                continue
            y = _parse_year_value(r[y_idx])
            if y is None:
                continue
            v = r[t_idx]
            if v is None:
                continue
            try:
                out[y] = float(v) / divisor
            except Exception:
                continue
        return out

    raise ValueError("Province target file must be .csv or .xlsx")


class CsvYieldDataset(Dataset):
    """
    Expects CSV with columns for date, county, target, and feature columns.
    Groups by (county, year) to build per-season samples.
    Optionally loads static phenological features (SOS, VGS, RGS).
    """

    def __init__(
        self,
        path: str,
        feature_cols: List[str],
        target_col: str = "Yield",
        date_col: str = "Date",
        county_col: str = "County",
        yield_csv: Optional[str] = None,
        yield_col: str = "yield",
        yield_county_col: str = "County",
        yield_year_col: str = "Year",
        weekly_csv: Optional[str] = None,
        weekly_cols: Optional[List[str]] = None,
        weekly_county_col: str = "name",
        weekly_year_col: str = "year",
        weekly_start_doy_col: str = "week_start_doy",
        weekly_end_doy_col: str = "week_end_doy",
        start_doy: int = 128,
        end_doy: int = 302,
        fill_missing: str = "ffill",
        static_cols: Optional[List[str]] = None,
        static_csv: Optional[str] = None,
        soil_csv: Optional[str] = None,
        soil_cols: Optional[List[str]] = None,
        weekly_mean: Optional[np.ndarray] = None,
        weekly_std: Optional[np.ndarray] = None,
        mean: Optional[np.ndarray] = None,
        std: Optional[np.ndarray] = None,
        y_mean: Optional[float] = None,
        y_std: Optional[float] = None,
        static_mean: Optional[np.ndarray] = None,
        static_std: Optional[np.ndarray] = None,
    ) -> None:
        self.mean = mean
        self.std = std
        self.y_mean = y_mean
        self.y_std = y_std
        self.weekly_mean = weekly_mean
        self.weekly_std = weekly_std
        self.static_mean = static_mean
        self.static_std = static_std
        self.feature_cols = feature_cols
        self.static_cols = static_cols or []
        self.hist_yield_col = "hist_yield_mean"
        self.hist_yield_missing_col = "hist_yield_missing_flag"
        self.static_cols_expanded = expand_static_cols(self.static_cols) if self.static_cols else []
        base_static_dim = len(self.static_cols_expanded)
        if self.static_cols:
            self.static_cols_expanded = self.static_cols_expanded + [
                self.hist_yield_col,
                self.hist_yield_missing_col,
            ]
        self.static_mask = None
        self.stats = {
            "groups_total": 0,
            "groups_with_target": 0,
            "samples_built": 0,
            "days_expected": 0,
            "days_present": 0,
            "static_missing": 0,
        }

        groups: Dict[Tuple[str, int], Dict[int, np.ndarray]] = {}
        targets: Dict[Tuple[str, int], float] = {}
        static_data: Dict[Tuple[str, int], np.ndarray] = {}

        if yield_csv:
            # 先读取独立产量表
            targets = _load_yield_table(
                yield_csv,
                yield_col=yield_col,
                yield_county_col=yield_county_col,
                yield_year_col=yield_year_col,
            )

        # 加载静态特征（如果提供）
        if static_csv or soil_csv:
            static_data = self._load_static_features(
                static_csv,
                soil_csv,
                static_cols,
                soil_cols or [],
                yield_county_col,
                yield_year_col,
            )

        weekly_cols = weekly_cols or []
        weekly_lookup = None
        weekly_starts_by_key = None
        if weekly_csv and weekly_cols:
            weekly_lookup, weekly_starts_by_key = self._load_weekly_features(
                weekly_csv,
                weekly_cols,
                weekly_county_col,
                weekly_year_col,
                weekly_start_doy_col,
                weekly_end_doy_col,
            )

        with open(path, "r", newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            headers = reader.fieldnames or []
            for col in [date_col, county_col] + feature_cols:
                if col not in headers:
                    raise ValueError(f"日数据缺少列: {col}")
            target_in_daily = target_col in headers

            for row in reader:
                dt = _parse_date(row[date_col])
                year = dt.year
                doy = dt.timetuple().tm_yday
                if doy < start_doy or doy > end_doy:
                    continue
                county = row[county_col]
                key = (county, year)

                feat = np.array([float(row[c]) for c in feature_cols], dtype=np.float32)
                groups.setdefault(key, {})[doy] = feat
                self.stats["days_present"] += 1

                if target_in_daily and row[target_col] != "":
                    # 允许日数据里也带产量
                    targets[key] = float(row[target_col])

        if not targets:
            raise ValueError("未找到产量数据，请提供产量列或 --yield-csv")

        county_history: Dict[str, List[Tuple[int, float]]] = {}
        for (county, year), yval in targets.items():
            county_history.setdefault(str(county), []).append((int(year), float(yval)))
        hist_mean_by_key: Dict[Tuple[str, int], float] = {}
        hist_missing_by_key: Dict[Tuple[str, int], float] = {}
        ema_alpha = 0.5
        for county, rows in county_history.items():
            rows = sorted(rows, key=lambda item: item[0])
            ema = None
            for year, yval in rows:
                if ema is None:
                    hist_mean_by_key[(county, year)] = 0.0
                    hist_missing_by_key[(county, year)] = 1.0
                    ema = float(yval)
                else:
                    hist_mean_by_key[(county, year)] = float(ema)
                    hist_missing_by_key[(county, year)] = 0.0
                    ema = ema_alpha * float(yval) + (1.0 - ema_alpha) * ema

        seq_len = end_doy - start_doy + 1
        self.stats["days_expected"] = seq_len
        x_list: List[np.ndarray] = []
        y_list: List[float] = []
        static_list: List[np.ndarray] = []
        static_mask_list: List[bool] = []
        static_dim = len(self.static_cols_expanded)
        weekly_dim = len(weekly_cols)
        weekly_list: List[np.ndarray] = []
        year_list: List[int] = []
        county_list: List[str] = []
        week_start_doys = [
            start_doy + i * 7
            for i in range(int(math.ceil(seq_len / 7)))
        ]
        expected_week_starts = set(week_start_doys)
        weekly_missing_idx = None
        if weekly_dim > 0:
            for idx, name in enumerate(weekly_cols):
                if "missing" in name.lower():
                    weekly_missing_idx = idx
                    break

        for key, by_day in groups.items():
            self.stats["groups_total"] += 1
            if key not in targets:
                continue
            self.stats["groups_with_target"] += 1
            county, year = key
            seq = np.full((seq_len, len(feature_cols)), np.nan, dtype=np.float32)
            for doy, feat in by_day.items():
                idx = doy - start_doy
                if 0 <= idx < seq_len:
                    seq[idx] = feat
            # 缺失日用填充策略补齐
            seq = _fill_missing(seq, fill_missing)
            x_list.append(seq)
            y_list.append(targets[key])
            year_list.append(int(year))
            county_list.append(str(county))

            if weekly_dim > 0:
                weekly_seq = np.zeros((len(week_start_doys), weekly_dim), dtype=np.float32)
                if weekly_lookup is not None:
                    if weekly_starts_by_key is not None:
                        starts = weekly_starts_by_key.get((str(county), int(year)), set())
                        if starts:
                            in_range = {d for d in starts if start_doy <= d <= end_doy}
                            if in_range and in_range != expected_week_starts:
                                sample = sorted(list(in_range))[:8]
                                raise ValueError(
                                    "Weekly SIF week_start_doy alignment mismatch "
                                    f"for {county}-{year}: expected {sorted(expected_week_starts)[:8]}..., "
                                    f"got {sample}..."
                                )
                    for w_idx, w_start in enumerate(week_start_doys):
                        weekly_row = weekly_lookup.get((str(county), int(year), int(w_start)))
                        if weekly_row is None:
                            if weekly_missing_idx is not None:
                                weekly_seq[w_idx, weekly_missing_idx] = 1.0
                            continue
                        for c_idx, col in enumerate(weekly_cols):
                            raw_val = weekly_row.get(col, "")
                            if raw_val == "":
                                if weekly_missing_idx is not None:
                                    weekly_seq[w_idx, weekly_missing_idx] = 1.0
                                continue
                            try:
                                val = float(raw_val)
                                if np.isnan(val):
                                    if weekly_missing_idx is not None:
                                        weekly_seq[w_idx, weekly_missing_idx] = 1.0
                                    continue
                                weekly_seq[w_idx, c_idx] = val
                            except ValueError:
                                if weekly_missing_idx is not None:
                                    weekly_seq[w_idx, weekly_missing_idx] = 1.0
                                continue
                if weekly_missing_idx is not None:
                    nan_rows = np.isnan(weekly_seq).any(axis=1)
                    weekly_seq[nan_rows, weekly_missing_idx] = 1.0
                weekly_seq = np.nan_to_num(weekly_seq, nan=0.0)
                weekly_list.append(weekly_seq)
            
            # 加载静态特征
            if static_dim > 0:
                hist_key = (str(county), int(year))
                hist_mean = hist_mean_by_key.get(hist_key, 0.0)
                hist_missing = hist_missing_by_key.get(hist_key, 1.0)
                if key in static_data:
                    base_static = static_data[key]
                    if base_static_dim > 0:
                        static_vec = np.concatenate(
                            [base_static, np.array([hist_mean, hist_missing], dtype=np.float32)]
                        )
                    else:
                        static_vec = np.array([hist_mean, hist_missing], dtype=np.float32)
                    static_list.append(static_vec)
                    static_mask_list.append(True)
                else:
                    static_vec = np.zeros(static_dim, dtype=np.float32)
                    static_vec[-2] = hist_mean
                    static_vec[-1] = hist_missing
                    static_list.append(static_vec)
                    static_mask_list.append(False)
                    self.stats["static_missing"] += 1

        if not x_list:
            raise ValueError("CSV 未构建出样本，请检查数据范围和列名")

        self.x = np.stack(x_list, axis=0)
        self.y = np.array(y_list, dtype=np.float32)
        self.static = np.stack(static_list, axis=0) if static_dim > 0 else None
        self.static_mask = (
            np.array(static_mask_list, dtype=bool) if static_dim > 0 else None
        )
        self.weekly = np.stack(weekly_list, axis=0) if weekly_dim > 0 else None
        self.years = np.array(year_list, dtype=np.int32)
        self.counties = np.array(county_list, dtype=object)
        self.x_t = None
        self.y_t = None
        self.static_t = None
        self.weekly_t = None
        self._apply_tensor_cache()

    @staticmethod
    def _load_static_features(
        path: str,
        soil_path: Optional[str],
        static_cols: List[str],
        soil_cols: List[str],
        county_col: str,
        year_col: str,
    ) -> Dict[Tuple[str, int], np.ndarray]:
        """从CSV加载静态特征数据，支持物候(按年) + 土壤(按县)"""
        static_data: Dict[Tuple[str, int], np.ndarray] = {}
        soil_cols = soil_cols or []
        soil_cols_set = set(soil_cols)

        if path and path.lower().endswith(".csv"):
            phenology_cols = [c for c in static_cols if c not in soil_cols_set]
            phenology_values: Dict[Tuple[str, int], Dict[str, str]] = {}
            with open(path, "r", newline="", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                headers = reader.fieldnames or []
                for col in [county_col, year_col] + phenology_cols:
                    if col not in headers:
                        raise ValueError(f"静态特征文件缺少列: {col}")
                for row in reader:
                    county = row[county_col]
                    year = int(row[year_col])
                    key = (str(county), int(year))
                    phenology_values[key] = row

            soil_values: Dict[str, Dict[str, str]] = {}
            if soil_path:
                with open(soil_path, "r", newline="", encoding="utf-8-sig") as f:
                    reader = csv.DictReader(f)
                    headers = reader.fieldnames or []
                    for col in [county_col] + soil_cols:
                        if col not in headers:
                            raise ValueError(f"土壤特征文件缺少列: {col}")
                    for row in reader:
                        county = row[county_col]
                        soil_values[str(county)] = row

            for key, row in phenology_values.items():
                county, year = key
                values_list: List[float] = []
                soil_row = soil_values.get(str(county), {})
                for col in static_cols:
                    if col in soil_cols_set:
                        raw_val = soil_row.get(col, "")
                    else:
                        raw_val = row.get(col, "")
                    if _is_peak_doy_col(col):
                        try:
                            doy = float(raw_val) if raw_val != "" else 0.0
                        except ValueError:
                            doy = 0.0
                        angle = 2.0 * math.pi * doy / 365.0
                        values_list.extend([math.sin(angle), math.cos(angle)])
                    else:
                        if raw_val == "":
                            values_list.append(0.0)
                        else:
                            try:
                                values_list.append(float(raw_val))
                            except ValueError:
                                values_list.append(0.0)
                static_data[(str(county), int(year))] = np.array(values_list, dtype=np.float32)
        return static_data

    @staticmethod
    def _load_weekly_features(
        path: str,
        weekly_cols: List[str],
        county_col: str,
        year_col: str,
        start_doy_col: str,
        end_doy_col: str,
    ) -> Tuple[Dict[Tuple[str, int, int], Dict[str, str]], Dict[Tuple[str, int], Set[int]]]:
        """Load weekly features keyed by (county, year, week_start_doy)."""
        lookup: Dict[Tuple[str, int, int], Dict[str, str]] = {}
        starts_by_key: Dict[Tuple[str, int], Set[int]] = {}
        with open(path, "r", newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            headers = reader.fieldnames or []
            for col in [county_col, year_col, start_doy_col, end_doy_col] + weekly_cols:
                if col not in headers:
                    raise ValueError(f"周数据缺少列: {col}")
            for row in reader:
                county = row[county_col]
                year = int(row[year_col])
                start_doy = int(row[start_doy_col])
                payload = {c: row.get(c, "") for c in weekly_cols}
                lookup[(str(county), year, start_doy)] = payload
                starts_by_key.setdefault((str(county), year), set()).add(start_doy)
        return lookup, starts_by_key

    def __len__(self) -> int:
        return self.x.shape[0]

    def __getitem__(self, idx: int) -> Tuple[torch.Tensor, ...]:
        if self.x_t is None or self.y_t is None:
            self._apply_tensor_cache()
        result = (self.x_t[idx], self.y_t[idx])
        if self.weekly_t is not None:
            result = result + (self.weekly_t[idx],)
        if self.static_t is not None:
            result = result + (self.static_t[idx],)
        return result

    def _apply_tensor_cache(self) -> None:
        x = self.x
        y = self.y
        static = self.static
        weekly = self.weekly
        if self.mean is not None and self.std is not None:
            x = (x - self.mean[None, None, :]) / self.std[None, None, :]
        if self.y_mean is not None and self.y_std is not None:
            y = (y - self.y_mean) / self.y_std
        if static is not None and self.static_mean is not None and self.static_std is not None:
            static = (static - self.static_mean[None, :]) / self.static_std[None, :]
        if weekly is not None and self.weekly_mean is not None and self.weekly_std is not None:
            weekly = (weekly - self.weekly_mean[None, None, :]) / self.weekly_std[None, None, :]
            self.weekly = weekly
        self.x_t = torch.from_numpy(x.astype(np.float32))
        self.y_t = torch.from_numpy(y.astype(np.float32)).view(-1)
        self.static_t = (
            torch.from_numpy(static.astype(np.float32)) if static is not None else None
        )
        self.weekly_t = (
            torch.from_numpy(weekly.astype(np.float32)) if weekly is not None else None
        )


class ArrayYieldDataset(Dataset):
    """
    直接使用内存中的 x/y 构建数据集，适合按年份切分后的训练/验证/测试。
    支持可选的静态特征输入。
    """

    def __init__(
        self,
        x: np.ndarray,
        y: np.ndarray,
        weekly: Optional[np.ndarray] = None,
        static: Optional[np.ndarray] = None,
        static_mask: Optional[np.ndarray] = None,
        years: Optional[np.ndarray] = None,
        mean: Optional[np.ndarray] = None,
        std: Optional[np.ndarray] = None,
        y_mean: Optional[float] = None,
        y_std: Optional[float] = None,
        weekly_mean: Optional[np.ndarray] = None,
        weekly_std: Optional[np.ndarray] = None,
        static_mean: Optional[np.ndarray] = None,
        static_std: Optional[np.ndarray] = None,
    ) -> None:
        x_np = x.astype(np.float32, copy=False)
        y_np = y.astype(np.float32, copy=False).reshape(-1)
        weekly_np = weekly.astype(np.float32, copy=False) if weekly is not None else None
        static_np = static.astype(np.float32, copy=False) if static is not None else None
        static_mask_np = static_mask.astype(bool, copy=False) if static_mask is not None else None
        self.years = years
        self.mean = mean
        self.std = std
        self.y_mean = y_mean
        self.y_std = y_std
        self.weekly_mean = weekly_mean
        self.weekly_std = weekly_std
        self.static_mean = static_mean
        self.static_std = static_std

        if self.mean is not None and self.std is not None:
            if x_np.ndim == 3:
                x_np = (x_np - self.mean[None, None, :]) / self.std[None, None, :]
            else:
                x_np = (x_np - self.mean[None, :]) / self.std[None, :]
        if self.y_mean is not None and self.y_std is not None:
            y_np = (y_np - float(self.y_mean)) / float(self.y_std)
        if weekly_np is not None and self.weekly_mean is not None and self.weekly_std is not None:
            weekly_np = (weekly_np - self.weekly_mean[None, None, :]) / self.weekly_std[None, None, :]
        if static_np is not None and self.static_mean is not None and self.static_std is not None:
            static_np = (static_np - self.static_mean[None, :]) / self.static_std[None, :]

        self.x_t = torch.from_numpy(x_np.astype(np.float32, copy=False))
        self.y_t = torch.from_numpy(y_np.astype(np.float32, copy=False)).view(-1)
        self.weekly_t = (
            torch.from_numpy(weekly_np.astype(np.float32, copy=False)) if weekly_np is not None else None
        )
        self.static_t = (
            torch.from_numpy(static_np.astype(np.float32, copy=False)) if static_np is not None else None
        )
        self.static_mask = torch.as_tensor(static_mask_np, dtype=torch.bool) if static_mask_np is not None else None
        
        if self.static_t is not None and self.static_t.shape[0] != self.x_t.shape[0]:
            raise ValueError("static 与 x 的样本数量不一致")
        if self.static_mask is not None and self.static_mask.shape[0] != self.x_t.shape[0]:
            raise ValueError("static_mask 与 x 的样本数量不一致")

    def __len__(self) -> int:
        return self.x_t.shape[0]

    def __getitem__(self, idx: int) -> Tuple[torch.Tensor, ...]:
        result = (self.x_t[idx], self.y_t[idx])
        if self.weekly_t is not None:
            result = result + (self.weekly_t[idx],)
        if self.static_t is not None:
            result = result + (self.static_t[idx],)
        return result


def compute_target_stats(y: np.ndarray, eps: float = 1e-6) -> Tuple[float, float]:
    mean = float(y.mean())
    std = float(y.std())
    if std < eps:
        std = 1.0
    return mean, std

def compute_static_stats(
    static: np.ndarray,
    mask: Optional[np.ndarray] = None,
    eps: float = 1e-6,
) -> Tuple[np.ndarray, np.ndarray]:
    """Compute static feature mean/std, ignoring missing rows when mask provided."""
    if static is None or static.size == 0:
        return np.zeros(0, dtype=np.float32), np.ones(0, dtype=np.float32)
    if mask is not None:
        mask = np.asarray(mask).astype(bool)
        if mask.any():
            static = static[mask]
        else:
            mean = np.zeros(static.shape[1], dtype=np.float32)
            std = np.ones(static.shape[1], dtype=np.float32)
            return mean, std
    mean = static.mean(axis=0)
    std = static.std(axis=0)
    std = np.where(std < eps, 1.0, std)
    return mean, std
